from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from sqlmodel import Session, select
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from ..database import get_session
from ..models import EngineSupply, GenSupply, SparePart  # ✅ التصحيح هنا
from typing import Dict, List, Optional
from ..models import (
    EngineSupply, EngineIssue, EngineRehab, EngineCheck, EngineUpload, EngineLathe, EnginePump, EngineElectrical,
    GenSupply, GenIssue, GenInspect, SparePart
)
from sqlmodel import select

router = APIRouter(prefix="/search", tags=["search"])

# ===== البحث =====
@router.get("/")
def global_search(q: str, session: Session = Depends(get_session)):
    q = q.strip()
    engines = session.exec(select(EngineSupply).where(EngineSupply.serial.contains(q))).all()
    gensets = session.exec(select(GenSupply).where(GenSupply.code.contains(q))).all()
    spares = session.exec(select(SparePart).where(SparePart.serial_or_code.contains(q))).all()  # ✅
    return {"engines": engines, "generators": gensets, "spares": spares}

# ===== التصدير =====
@router.get("/export")
def export_excel(q: str, session: Session = Depends(get_session)):
    q = q.strip()
    wb = Workbook()
    ws = wb.active
    ws.title = "نتائج البحث"
    ws.sheet_view.rightToLeft = True

    header_fill = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
    ws.append(["النوع", "الرقم", "التفاصيل", "الملاحظات"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        cell.fill = header_fill

    engines = session.exec(select(EngineSupply).where(EngineSupply.serial.contains(q))).all()
    for e in engines:
        ws.append(["محرك", e.serial, f"{e.engine_type or ''} - {e.model or ''}", e.notes or ""])

    gensets = session.exec(select(GenSupply).where(GenSupply.code.contains(q))).all()
    for g in gensets:
        ws.append(["مولد", g.code, f"{g.gen_type or ''} - {g.model or ''}", g.notes or ""])

    spares = session.exec(select(SparePart).where(SparePart.serial_or_code.contains(q))).all()  # ✅
    for s in spares:
        ws.append(["قطعة", s.serial_or_code, f"{s.part_name} × {s.qty}", s.notes or ""])

    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 25

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="search_{q}.xlsx"'}
    )
# ===== ملخّص مجمّع لكل رقم/ترميز عبر جميع العمليات =====
def _latest(rows: List, date_field: str = "date"):
    rows = [r for r in rows if getattr(r, date_field, None)]
    return sorted(rows, key=lambda r: getattr(r, date_field))[-1] if rows else None

@router.get("/summary")
def search_summary(q: str, session: Session = Depends(get_session)):
    q = (q or "").strip()

    # --- اجمع كل السجلات ذات الصلة ---
    e_sup  = session.exec(select(EngineSupply).where(EngineSupply.serial.contains(q))).all()
    e_iss  = session.exec(select(EngineIssue).where(EngineIssue.serial.contains(q))).all()
    e_reh  = session.exec(select(EngineRehab).where(EngineRehab.serial.contains(q))).all()
    e_chk  = session.exec(select(EngineCheck).where(EngineCheck.serial.contains(q))).all()
    e_up   = session.exec(select(EngineUpload).where(EngineUpload.serial.contains(q))).all()
    e_lat  = session.exec(select(EngineLathe).where(EngineLathe.serial.contains(q))).all()
    e_pmp  = session.exec(select(EnginePump).where(EnginePump.serial.contains(q))).all()
    e_elec = session.exec(select(EngineElectrical).where(EngineElectrical.serial.contains(q))).all()

    g_sup  = session.exec(select(GenSupply).where(GenSupply.code.contains(q))).all()
    g_iss  = session.exec(select(GenIssue).where(GenIssue.code.contains(q))).all()
    g_ins  = session.exec(select(GenInspect).where(GenInspect.code.contains(q))).all()

    spares = session.exec(select(SparePart).where(SparePart.serial_or_code.contains(q))).all()

    # --- جمّع حسب المفتاح (serial أو code) ---
    engines: Dict[str, Dict[str, List]] = {}
    for r in e_sup:  engines.setdefault(r.serial, {"sup":[], "iss":[], "reh":[], "chk":[], "up":[], "lat":[], "pmp":[], "elec":[], "sp": []})["sup"].append(r)
    for r in e_iss:  engines.setdefault(r.serial, {"sup":[], "iss":[], "reh":[], "chk":[], "up":[], "lat":[], "pmp":[], "elec":[], "sp": []})["iss"].append(r)
    for r in e_reh:  engines.setdefault(r.serial, {"sup":[], "iss":[], "reh":[], "chk":[], "up":[], "lat":[], "pmp":[], "elec":[], "sp": []})["reh"].append(r)
    for r in e_chk:  engines.setdefault(r.serial, {"sup":[], "iss":[], "reh":[], "chk":[], "up":[], "lat":[], "pmp":[], "elec":[], "sp": []})["chk"].append(r)
    for r in e_up:   engines.setdefault(r.serial, {"sup":[], "iss":[], "reh":[], "chk":[], "up":[], "lat":[], "pmp":[], "elec":[], "sp": []})["up"].append(r)
    for r in e_lat:  engines.setdefault(r.serial, {"sup":[], "iss":[], "reh":[], "chk":[], "up":[], "lat":[], "pmp":[], "elec":[], "sp": []})["lat"].append(r)
    for r in e_pmp:  engines.setdefault(r.serial, {"sup":[], "iss":[], "reh":[], "chk":[], "up":[], "lat":[], "pmp":[], "elec":[], "sp": []})["pmp"].append(r)
    for r in e_elec: engines.setdefault(r.serial, {"sup":[], "iss":[], "reh":[], "chk":[], "up":[], "lat":[], "pmp":[], "elec":[], "sp": []})["elec"].append(r)

    gensets: Dict[str, Dict[str, List]] = {}
    for r in g_sup:  gensets.setdefault(r.code, {"sup":[], "iss":[], "ins":[], "sp": []})["sup"].append(r)
    for r in g_iss:  gensets.setdefault(r.code, {"sup":[], "iss":[], "ins":[], "sp": []})["iss"].append(r)
    for r in g_ins:  gensets.setdefault(r.code, {"sup":[], "iss":[], "ins":[], "sp": []})["ins"].append(r)

    # اربط قطع الغيار حسب المفتاح (serial_or_code)
    for s in spares:
        k = s.serial_or_code
        if k in engines: engines[k]["sp"].append(s)
        if k in gensets: gensets[k]["sp"].append(s)

    # --- صفوف ملخصة (آخر سجل لكل عملية) ---
    rows: List[Dict] = []

    for serial, g in engines.items():
        sup, iss, reh, chk, up, lat, pmp, elec = map(_latest, (g["sup"], g["iss"], g["reh"], g["chk"], g["up"], g["lat"], g["pmp"], g["elec"]))
        last_sp = _latest(g["sp"], "id")  # مجرد آخر إدخال قطع غيار (لا يحتوي تاريخ إلزامي)
        rows.append({
            "kind": "محرك",
            "key": serial,
            # توريد
            "supply_type": getattr(sup, "engine_type", None),
            "supply_model": getattr(sup, "model", None),
            "supply_supplier": getattr(sup, "supplier", None),
            "supply_prev_site": getattr(sup, "prev_site", None),
            "supply_date": getattr(sup, "date", None),
            # صرف
            "issue_current_site": getattr(iss, "current_site", None),
            "issue_receiver": getattr(iss, "receiver", None),
            "issue_requester": getattr(iss, "requester", None),
            "issue_date": getattr(iss, "date", None),
            # تأهيل
            "rehab_by": getattr(reh, "rehab_by", None),
            "rehab_type": getattr(reh, "rehab_type", None),
            "rehab_date": getattr(reh, "date", None),
            # فحص
            "check_inspector": getattr(chk, "inspector", None),
            "check_desc": getattr(chk, "description", None),
            "check_date": getattr(chk, "date", None),
            # رفع
            "upload_rehab_file": getattr(up, "rehab_file", None),
            "upload_check_file": getattr(up, "check_file", None),
            "upload_rehab_date": getattr(up, "rehab_date", None),
            "upload_check_date": getattr(up, "check_date", None),
            # مخرطة
            "lathe_rehab": getattr(lat, "lathe_rehab", None),
            "lathe_supply_date": getattr(lat, "lathe_supply_date", None),
            # بمبات
            "pump_serial": getattr(pmp, "pump_serial", None),
            "pump_rehab": getattr(pmp, "pump_rehab", None),
            # صريمي
            "elect_kind": getattr(elec, "kind", None),
            "elect_starter": getattr(elec, "has_starter", None),
            "elect_dynamo": getattr(elec, "has_dynamo", None),
            # قطع غيار
            "spare_last": f"{getattr(last_sp,'part_name', '')} × {getattr(last_sp,'qty', '')}" if last_sp else None
        })

    for code, g in gensets.items():
        sup, iss, ins = _latest(g["sup"]), _latest(g["iss"]), _latest(g["ins"])
        last_sp = _latest(g["sp"], "id")
        rows.append({
            "kind": "مولد",
            "key": code,
            # توريد
            "supply_type": getattr(sup, "gen_type", None),
            "supply_model": getattr(sup, "model", None),
            "supply_supplier": getattr(sup, "supplier_name", None),
            "supply_entity": getattr(sup, "supplier_entity", None),
            "supply_prev_site": getattr(sup, "prev_site", None),
            "supply_date": getattr(sup, "date", None),
            # صرف
            "issue_date": getattr(iss, "issue_date", None),
            "issue_receiver": getattr(iss, "receiver", None),
            "issue_requester": getattr(iss, "requester", None),
            "issue_current_site": getattr(iss, "current_site", None),
            # فحص/رفع
            "inspect_inspector": getattr(ins, "inspector", None),
            "inspect_elect_by": getattr(ins, "electrical_rehab_by", None),
            "inspect_rehab_date": getattr(ins, "rehab_date", None),
            # قطع غيار
            "spare_last": f"{getattr(last_sp,'part_name', '')} × {getattr(last_sp,'qty', '')}" if last_sp else None
        })

    # فرز بسيط بالنوع ثم المفتاح
    rows.sort(key=lambda r: (r["kind"], str(r["key"])))
    return {"rows": rows}
